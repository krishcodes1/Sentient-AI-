"""Reads Excel (.xlsx) workbooks with openpyxl in read-only, data-only mode.

Why it exists: gradebooks and budgets arrive as spreadsheets. data_only means
the cached values are read and no formula is ever evaluated; read_only
streams rows. The archive is checked against the shared zip limits first
(ooxml.safe_zip), then at most 20 sheets, 5000 rows per sheet, 60 columns
and 1000 characters per cell are read. Rows become tab-separated lines,
grouped into units of about 3000 characters labelled "Sheet 'Grades' rows
1-120" (the page number is the sheet's position).
"""

from __future__ import annotations

import datetime as dt
import io
from typing import Any

from services.files.limits import (
    MAX_CELL_CHARS,
    MAX_XLSX_COLUMNS,
    MAX_XLSX_ROWS,
    MAX_XLSX_SHEETS,
    SECTION_TARGET_CHARS,
)
from services.files.worker.parsers.ooxml import safe_zip
from services.files.worker.protocol import Emitter, ParseError


def cell_text(value: Any) -> str:
    """One cell as text, capped at MAX_CELL_CHARS."""
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, float):
        text = format(value, ".15g")
    elif isinstance(value, (dt.datetime, dt.date, dt.time)):
        text = value.isoformat()
    else:
        text = str(value)
    text = " ".join(text.replace("\t", " ").split())
    return text[:MAX_CELL_CHARS]


def _sheet_label(name: str, first: int, last: int) -> str:
    safe = " ".join(str(name).split())[:60].replace("'", "’")
    return f"Sheet '{safe}' rows {first}-{last}"


def parse(data: bytes, header: dict[str, Any], emit: Emitter) -> None:
    safe_zip(data).close()
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many types for bad files
        raise ParseError("corrupt") from None
    try:
        sheets = list(workbook.worksheets)
        emit.meta(title="", pages_total=len(sheets))
        truncated = len(sheets) > MAX_XLSX_SHEETS
        for number, sheet in enumerate(sheets[:MAX_XLSX_SHEETS], start=1):
            name = getattr(sheet, "title", f"Sheet{number}")
            lines: list[str] = []
            size = 0
            first = 0
            last = 0
            rows = sheet.iter_rows(
                min_row=1, max_row=MAX_XLSX_ROWS + 1, max_col=MAX_XLSX_COLUMNS, values_only=True
            )
            for row_number, row in enumerate(rows, start=1):
                if row_number > MAX_XLSX_ROWS:
                    truncated = True
                    break
                cells = [cell_text(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                line = "\t".join(cells)
                if lines and size + len(line) > SECTION_TARGET_CHARS:
                    emit.unit(_sheet_label(name, first, last), number, "\n".join(lines))
                    lines, size = [], 0
                if not lines:
                    first = row_number
                lines.append(line)
                size += len(line) + 1
                last = row_number
            if lines:
                emit.unit(_sheet_label(name, first, last), number, "\n".join(lines))
        emit.done(truncated=truncated)
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass
